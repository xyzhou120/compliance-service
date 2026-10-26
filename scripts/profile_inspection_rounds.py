"""Profile repeatable local document checks without model or callback traffic."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import statistics
import sys
import shutil
from time import perf_counter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from compliance_service.config import load_model_settings
from compliance_service.ledger import Ledger
from compliance_service.service import InspectionService
from compliance_service.sources import LocalDocumentSource
from openpyxl import load_workbook


CASES = (
    ("complete-cold", "full-check-samples/complete", "stop_work_unload"),
    ("complete-warm-1", "full-check-samples/complete", "stop_work_unload"),
    ("complete-warm-2", "full-check-samples/complete", "stop_work_unload"),
    ("bad-relief", "full-check-samples/bad-relief", "stop_work_unload"),
    ("missing-signature", "full-check-samples/missing-signature", "stop_work_unload"),
    ("stop-not-immediate", "generated/stop-not-immediate", "stop_work_unload"),
    ("power-off-not-immediate", "generated/power-off-not-immediate", "stop_work_unload"),
    ("evacuation-route-wrong", "generated/evacuation-route-wrong", "stop_work_unload"),
    ("bad-effect", "full-check-samples/bad-effect", "relief_effects"),
    ("missing-dispatch", "document-cases/missing-dispatch", "relief_effects"),
)


def run_rounds(root: Path, output: Path) -> dict:
    root = root.resolve()
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    for name, action, column in (
        ("stop-not-immediate", "停工", "立即"),
        ("power-off-not-immediate", "断电", "立即"),
        ("evacuation-route-wrong", "撤人", "路线符合"),
    ):
        destination = output / "generated" / name
        destination.mkdir(parents=True, exist_ok=True)
        original = root / "full-check-samples" / "complete"
        for source_file in original.iterdir():
            if source_file.is_file():
                shutil.copy2(source_file, destination / source_file.name)
        workbook = load_workbook(destination / "02-dispatch.xlsx")
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            if row[headers.index("动作")].value == action:
                row[headers.index(column)].value = "否"
                break
        workbook.save(destination / "02-dispatch.xlsx")
    settings = load_model_settings()
    settings["ocr"]["enabled"] = False
    settings["vision"]["enabled"] = False
    ledger = Ledger(output / "rounds.sqlite3")
    source = LocalDocumentSource(output / "artifacts", model_settings=settings)

    def source_with_progress(run):
        def progress(event):
            ledger.record_progress(run["run_id"], run["worker_id"], run["fence"], event)
        return source(run, progress_callback=progress)

    service = InspectionService(ledger, source=source_with_progress)
    rows = []
    for name, folder, item_code in CASES:
        path = (output if folder.startswith("generated/") else root) / folder / "manifest.json"
        manifest = json.loads(path.read_text(encoding="utf-8-sig"))
        accepted = service.submit({
            "sourceSystem": manifest["source_system"], "mineCode": manifest["mine_code"],
            "environment": manifest["environment"], "checkCycle": "2026-10",
            "scope": {"manifest": str(path)}, "callbackMode": "none",
            "items": [{"itemCode": item_code}],
        })["items"][0]
        started = perf_counter()
        service.run_once(run_id=accepted["run_id"])
        total_ms = round((perf_counter() - started) * 1000, 2)
        run = ledger.get(accepted["run_id"], manifest["source_system"], manifest["mine_code"], manifest["environment"])
        result, snapshot = run["result"], run["snapshot"]
        progress = ledger.get_progress(accepted["run_id"], manifest["source_system"],
                                       manifest["mine_code"], manifest["environment"], limit=500)
        stages = {stage: round(sum(event.get("elapsed_ms", 0) for event in progress["events"]
                                   if event.get("stage") == stage and event.get("event") ==
                                   ("extraction_finished" if stage == "extraction" else "finished")), 2)
                  for stage in ("source", "extraction", "evaluation")}
        rows.append({
            "case": name, "origin": "synthetic_document", "manifest": str(path),
            "run_id": accepted["run_id"], "item_code": item_code,
            "status": result["status"], "execution_state": run["execution_state"],
            "total_ms": total_ms, "stage_ms": stages,
            "file_count": len(snapshot["files"]) if snapshot else 0,
            "evidence_count": len(snapshot["evidence"]) if snapshot else 0,
            "cache": snapshot.get("cache", {}) if snapshot else {},
            "finding_codes": [finding.get("code") for finding in result.get("findings", [])],
            "findings": result.get("findings", []),
            "coverage_gap_count": len(result.get("coverage_gaps", [])),
            "description": result["description"],
            "callback_state": run["delivery_state"],
        })
        print(json.dumps({key: rows[-1][key] for key in
                          ("case", "item_code", "status", "total_ms", "stage_ms", "cache", "finding_codes")},
                         ensure_ascii=False), flush=True)
    warm = [row["total_ms"] for row in rows if row["case"].startswith("complete-warm")]
    report = {
        "profile": "local-native-documents-no-model-no-callback",
        "rounds": rows,
        "complete_cold_ms": rows[0]["total_ms"],
        "complete_warm_median_ms": round(statistics.median(warm), 2),
        "warnings": ["Synthetic cases test engineering behavior, not real mine accuracy.",
                     "OCR/VLM disabled for repeatable timing; scanned-document latency is excluded."],
    }
    (output / "profile.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def profile_real_native(manifest_path: Path, output: Path) -> dict:
    """Measure the same real bytes with recognition disabled as a flow baseline."""
    manifest_path = manifest_path.resolve()
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    settings = load_model_settings()
    settings["ocr"]["enabled"] = False
    settings["vision"]["enabled"] = False
    ledger = Ledger(output / "real-native.sqlite3")
    source = LocalDocumentSource(output / "real-native-artifacts", model_settings=settings)

    def source_with_progress(run):
        return source(run, progress_callback=lambda event: ledger.record_progress(
            run["run_id"], run["worker_id"], run["fence"], event))

    service = InspectionService(ledger, source=source_with_progress)
    accepted = service.submit({
        "sourceSystem": manifest["source_system"], "mineCode": manifest["mine_code"],
        "environment": manifest["environment"], "checkCycle": "2026-10",
        "scope": {"manifest": str(manifest_path)}, "callbackMode": "none",
        "items": [{"itemCode": "stop_work_unload"}],
    })["items"][0]
    started = perf_counter()
    service.run_once(run_id=accepted["run_id"])
    elapsed_ms = round((perf_counter() - started) * 1000, 2)
    run = ledger.get(accepted["run_id"], manifest["source_system"], manifest["mine_code"], manifest["environment"])
    progress = ledger.get_progress(accepted["run_id"], manifest["source_system"],
                                   manifest["mine_code"], manifest["environment"], limit=500)
    source_ms = round(sum(event.get("elapsed_ms", 0) for event in progress["events"]
                          if event.get("stage") == "source" and event.get("event") == "finished"), 2)
    outcome = {
        "origin": "real_documents_flow_control_no_ocr_vlm", "manifest": str(manifest_path),
        "run_id": accepted["run_id"], "status": run["result"]["status"], "total_ms": elapsed_ms,
        "source_ms": source_ms, "file_count": len(run["snapshot"]["files"]),
        "coverage_gap_count": len(run["result"].get("coverage_gaps", [])),
        "cache": run["snapshot"].get("cache", {}), "callback_state": run["delivery_state"],
        "note": "Recognition disabled; this result is only a timing/flow comparison.",
    }
    (output / "real-native-profile.json").write_text(json.dumps(outcome, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(outcome, ensure_ascii=False), flush=True)
    return outcome


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-root", type=Path, default=Path(".local/validation"))
    parser.add_argument("--output", type=Path, default=Path(".local/validation/performance-20261026"))
    parser.add_argument("--real-native-manifest", type=Path)
    parser.add_argument("--real-only", action="store_true")
    args = parser.parse_args()
    if not args.real_only:
        run_rounds(args.input_root, args.output)
    if args.real_native_manifest:
        profile_real_native(args.real_native_manifest, args.output)
