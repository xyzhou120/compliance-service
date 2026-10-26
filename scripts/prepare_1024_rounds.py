from __future__ import annotations

"""Select small, repeatable debugging rounds from the 1024.xlsx inventory.

This script is intentionally read-only against the source workbook. It de-duplicates
attachment IDs, keeps same-name/different-ID records separate, and emits request
payloads that preserve source row and document metadata for the service layer.
"""

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

GROUP_MAP = {
    "防冲工作日报表": "prevention_daily",
    "调度日报": "dispatch_daily",
    "生产矿长（总工程师）日分析记录": "analysis",
    "防冲钻孔验收记录": "relief_acceptance",
    "调度解除": "release",
}

LOCAL_1014_16 = {
    # These files are the already downloaded 10/14-10/16 bundle. The first
    # three files in that directory are the design/indicator documents, so
    # the daily records start at 04-dispatch_daily.xlsx.
    "2110534812013506561": "04-dispatch_daily.xlsx",
    "2110620560750358530": "05-prevention_daily.pdf",
    "2110620731785687041": "06-analysis.pdf",
    "2110900926683926530": "07-dispatch_daily.xlsx",
    "2110901418180857858": "08-dispatch_daily.xlsx",
    "2110899702035238914": "09-prevention_daily.pdf",
    "2110899839713267714": "10-analysis.pdf",
    "2111261288597893122": "11-dispatch_daily.xlsx",
    "2111256851315806210": "12-prevention_daily.pdf",
    "2111256982914150402": "13-analysis.pdf",
}


def parse_day(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text or text.endswith("/00"):
        return None
    for fmt in ("%Y/%m/%d", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    m = re.search(r"(20\d{2})[/-](\d{1,2})[/-](\d{1,2})", text)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def canonical_group(name: str) -> str | None:
    name = str(name or "").strip()
    for prefix, group in GROUP_MAP.items():
        if prefix in name:
            return group
    return None


def load_inventory(source: Path) -> tuple[list[dict[str, Any]], str]:
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb["清单和煤矿常用信息查询"]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    by_id: dict[str, dict[str, Any]] = {}
    for source_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        row = dict(zip(header, values))
        aid = str(row.get("attachment_id") or "").strip()
        if not aid or aid in by_id:
            continue
        group = canonical_group(row.get("file_name"))
        if not group:
            continue
        day = parse_day(row.get("report_time"))
        if day is None:
            continue
        by_id[aid] = {
            "mine_code": str(row.get("mine_code") or ""),
            "mine_name": str(row.get("mine_name") or ""),
            "report_date": day.isoformat(),
            "report_time": str(row.get("report_time") or ""),
            "update_rate": row.get("update_rate"),
            "cycle": row.get("cycle"),
            "file_name_id": row.get("file_name_id"),
            "file_name": row.get("file_name"),
            "attachment_id": aid,
            "attachment_name": row.get("attachment_name"),
            "attachment_link": row.get("attachment_link"),
            "suffix": row.get("suffix"),
            "source_row": source_row,
            "material_group": group,
        }
    return list(by_id.values()), sha


def within(rows: list[dict[str, Any]], start: date, end: date) -> list[dict[str, Any]]:
    return [r for r in rows if start.isoformat() <= r["report_date"] <= end.isoformat()]


def select_round(rows: list[dict[str, Any]], start: date, end: date, *, include_acceptance: bool = False,
                 include_version_conflicts: bool = False) -> list[dict[str, Any]]:
    candidates = within(rows, start, end)
    chosen: list[dict[str, Any]] = []
    # Prefer the first record per date/group; retain distinct IDs for explicit conflicts.
    by_day_group: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in candidates:
        by_day_group[(row["report_date"], row["material_group"])].append(row)
    for day in (start + timedelta(days=i) for i in range((end - start).days + 1)):
        day_s = day.isoformat()
        for group in ("prevention_daily", "dispatch_daily", "analysis"):
            items = by_day_group.get((day_s, group), [])
            if not items:
                continue
            chosen.append(sorted(items, key=lambda r: (r["source_row"], r["attachment_id"]))[0])
            if include_version_conflicts and group in {"prevention_daily", "dispatch_daily"}:
                # Keep one extra distinct attachment ID to make version handling observable.
                for item in sorted(items, key=lambda r: (r["source_row"], r["attachment_id"]))[1:]:
                    chosen.append(item)
                    break
    if include_acceptance:
        acceptance = [r for r in rows if r["material_group"] == "relief_acceptance" and r["report_date"] <= end.isoformat()]
        if acceptance:
            chosen.append(sorted(acceptance, key=lambda r: (r["report_date"], r["source_row"]))[-1])
    return chosen


def source_metadata(row: dict[str, Any], source: Path) -> dict[str, Any]:
    return {
        "attachment_id": row["attachment_id"],
        "report_date": row["report_date"],
        "report_time": row["report_time"],
        "document_type": row["file_name"],
        "material_group": row["material_group"],
        "source_row": row["source_row"],
        "file_name_id": row["file_name_id"],
        "attachment_name": row["attachment_name"],
        "source_workbook": str(source),
    }


def request_for(items: list[dict[str, Any]], source: Path, start: date, end: date,
                *, local_dir: Path | None = None) -> dict[str, Any]:
    mine_code = items[0]["mine_code"]
    attachments: list[dict[str, Any]] = []
    for row in items:
        aid = row["attachment_id"]
        local_name = LOCAL_1014_16.get(aid) if local_dir else None
        item: dict[str, Any] = {
            "group_id": row["material_group"],
            "name": row["attachment_name"],
            "source_metadata": source_metadata(row, source),
        }
        if local_name:
            local_path = (local_dir / local_name).resolve()
            if not local_path.is_file():
                raise FileNotFoundError(f"Previously downloaded attachment is missing: {local_path}")
            item.update({"kind": "local_path", "path": str(local_path)})
        else:
            item.update({"kind": "url", "url": row["attachment_link"]})
        attachments.append(item)
    return {
        "sourceSystem": "excel-1024-debug",
        "mineCode": mine_code,
        "environment": "test",
        "checkCycle": start.strftime("%Y-%m"),
        "ruleVersion": "impact-1023-draft-v1",
        "items": [{"itemCode": "stop_work_unload"}],
        "period": {"start": start.isoformat(), "end": end.isoformat()},
        "source_complete": False,
        "attachments": attachments,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path(r"C:\Users\86158\Desktop\1024.xlsx"))
    parser.add_argument("--out", type=Path, default=Path(".local/validation/1024-multi-round"))
    parser.add_argument("--local-dir", type=Path, default=Path(".local/validation/hujiahe-1024-20261014-16/source"))
    args = parser.parse_args()
    rows, sha = load_inventory(args.source)
    args.out.mkdir(parents=True, exist_ok=True)
    rounds = [
        ("round-1014-1016-existing", date(2026, 10, 14), date(2026, 10, 16), False, True, True),
        ("round-0923-0925-effective", date(2026, 9, 23), date(2026, 9, 25), False, False, True),
        ("round-1020-1022-latest", date(2026, 10, 20), date(2026, 10, 22), True, False, False),
    ]
    # Keep a small audit summary in the generated selection file so the
    # parent runner can explain why a record was selected without reopening
    # the workbook. Duplicate attachment IDs were removed before selection;
    # same-name records with different IDs remain distinct by design.
    wb = load_workbook(args.source, read_only=True, data_only=True)
    ws = wb["清单和煤矿常用信息查询"]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    all_rows = [dict(zip(header, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
    id_counts = defaultdict(int)
    name_to_ids: dict[str, set[str]] = defaultdict(set)
    for row in all_rows:
        aid = str(row.get("attachment_id") or "").strip()
        if aid:
            id_counts[aid] += 1
            name_to_ids[str(row.get("attachment_name") or "")].add(aid)
    same_name_distinct_ids = [
        {"attachment_name": name, "attachment_ids": sorted(ids)}
        for name, ids in sorted(name_to_ids.items()) if name and len(ids) > 1
    ]
    manifest: dict[str, Any] = {
        "source_workbook": str(args.source),
        "source_sha256": sha,
        "sheet": "清单和煤矿常用信息查询",
        "inventory_rows": len(all_rows),
        "unique_attachment_ids": len(id_counts),
        "duplicate_attachment_rows_removed": sum(n - 1 for n in id_counts.values() if n > 1),
        "same_name_distinct_id_examples": same_name_distinct_ids[:50],
        "rounds": [],
    }
    for name, start, end, include_acceptance, use_local, conflicts in rounds:
        items = select_round(rows, start, end, include_acceptance=include_acceptance, include_version_conflicts=conflicts)
        local_dir = args.local_dir if use_local else None
        payload = request_for(items, args.source, start, end, local_dir=local_dir)
        path = args.out / f"{name}.request.json"
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        conflict_groups = []
        grouped = defaultdict(list)
        for row in items:
            grouped[(row["report_date"], row["material_group"])].append(row["attachment_id"])
        for key, ids in grouped.items():
            if len(ids) > 1:
                conflict_groups.append({"report_date": key[0], "material_group": key[1], "attachment_ids": ids})
        rationale = {
            "round-1014-1016-existing": "复用已下载的10/14-10/16轻量日报/日分析；保留10/15两个调度日报附件，观察版本冲突。",
            "round-0923-0925-effective": "选取2026-09-23细则生效后的连续三日，保留同日不同attachment_id版本，验证重复版本处理。",
            "round-1020-1022-latest": "选择最新完整的10/20-10/22连续三日；加入一份截至10/18的周钻孔验收作负向关联样例，不能凭该记录证明10/20-10/22卸压已完成。",
        }[name]
        manifest["rounds"].append({
            "name": name, "period": {"start": start.isoformat(), "end": end.isoformat()},
            "attachment_count": len(items), "request": str(path), "attachment_ids": [r["attachment_id"] for r in items],
            "material_groups": dict(sorted(Counter(r["material_group"] for r in items).items())),
            "version_conflicts": conflict_groups,
            "source_rows": [r["source_row"] for r in items],
            "selection_rationale": rationale,
        })
    (args.out / "selection.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
