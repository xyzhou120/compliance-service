from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
OUT.mkdir(exist_ok=True)
MD_PATH = OUT / "冲击地压跨文档合规检查服务开发交接与验收方案.md"
DOCX_PATH = OUT / "冲击地压跨文档合规检查服务开发交接与验收方案.docx"

MD = r'''# 冲击地压跨文档合规检查服务开发交接与验收方案

版本：开发交接版 2026-10-25  
面向对象：服务开发、数据接入、规则实现、测试和联调人员

## 一 结论和交接边界

本项目的目标不是增加一个本地文件调试器，而是建设一条生产主链：

```text
上游任务下发
  -> 按 mineId、checkCycle、items 查询 OceanBase 清单
  -> 冻结本轮附件清单和文件字节版本
  -> 解析防冲日报、调度日报、日分析、设计、临界指标、验收等多类文件
  -> 识别防冲日报预警或动力线索
  -> 关联撤人、停工、断电、报告、卸压和验收记录
  -> 形成逐日问题和可解释的连续事件
  -> 每个检查项独立回调上游
```

当前代码已经完成单机本地调试链路、原生文档解析、跨文件事实与证据保存、两项规则原型、OCR/VLM 候选接入、日报线索和周期聚合纯函数、文件/页面/provider 进度记录以及调试页面。当前代码仍未完成真实生产主链：上游任务协议、OceanBase 查询、正式附件取证、真实结果回调、暂停/恢复和生产身份授权均未接入。

当前不能把 `execution_state=completed` 当作业务核验正确。它只表示执行流程收束，pending 和 unverifiable 也可以是本轮终态。真实业务结论必须同时满足事实、证据定位、事件归属、规则覆盖和资料完整性条件。

本服务内不提供人工修改 pass/no_pass 的入口。下游核实反馈用于优化能力和触发新运行，历史结果不覆盖。分层盲测保留为 M6 最终验收，不提前变成日常审批。跨矿、跨运行证据隔离保留为必要底线；SSRF、provider 外发和回调目标约束在对应真实网络接入时落实。本次只编制开发交接文件，未修改服务逻辑或连接数据库。

## 二 用户确认的两项检查和规则基线

检查项来自 `冲击地压规则1023.xlsx` 的 `AI深度监察` Sheet 第 2、3 行：

| itemCode | 名称 | 业务主线 |
| --- | --- | --- |
| `stop_work_unload` | 冲击危险停工撤人与卸压 | 预警或动力线索 -> 研判 -> 停工、断电、撤人、报告 -> 卸压措施和验收 -> 继续采掘反证 |
| `relief_effects` | 冲击地压解危效果检验 | 解危施工 -> 效果检验原始数据 -> 阈值比较 -> 技术负责人确认 -> 调度解除 -> 实际复工 |

当前代码登记的规则源哈希为 `6c0fe47c...`，本次重新读取的 `D:\西安院\深度监察\冲击地压\冲击地压规则1023.xlsx` 哈希为 `e3fc04ca...`。两者不一致，开发必须先确认文件是否换版，并登记新的不可变 `rulebook_revision`。不能用同一个版本名代表两个不同文件内容。

规则 Excel 的 `J2/J3` 是原表参考结果，不能作为真实矿井结果输入。规则说明明确要求：目录有清单不等于行为合规；查无材料不等于未实施；缺证、时间不清或冲突应进入待补证/人工复核；等于阈值等边界情况不得自动当作通过。

本次逐单元格比较代码保存的 B2:J3，确认内容差异在 F2。当前 F2 增加分层判定，强调“仅证实未断电或未报告调度”不能单独套用重大隐患条款，且查无材料不等于未实施。当前文件 SHA256 为 `e3fc04ca7346bb91bd8868b24efcd42155df1ae68c2a41a5f01a63a0da1e1247`；旧登记源 SHA256 为 `6c0fe47c266d5af9eeae41953c7dc1fc4b685620ee8fdcbc2e96d3d14c27f9fa`。

## 三 上游任务和结果回调契约

### 3.1 任务下发

上游任务示例要求检查服务接收：

```json
{
  "mineId": "610203006642",
  "status": "running",
  "checkCycle": "2026-10",
  "items": [
    {"itemCode": "GAS_DAILY_001", "itemName": "瓦斯日报检查"}
  ],
  "resultCallbackUrl": "https://api.example.com/fileAudit/mineDataInspectionTask/resultCallback"
}
```

响应成功只表示控制动作已生效，不表示检查结果已经产生。每个检查项独立执行、独立回调。服务必须保存任务原文、任务 ID（若上游提供）、矿井、周期、检查项列表、状态和回调 URL 快照。

当前代码只有 `POST /v1/manual/runs` 和 `POST /v1/upstream/runs`，使用 `sourceSystem`、`mineCode`、`scope.manifest` 和 `callbackMode`。真实上游请求目前不能直接被当前合同接受。

正式入口由基础表 `check_url` 决定，样例为 `POST /inspection/tasks`。不要自行用 `/v1/upstream/tasks` 替换上游已配置地址。成功响应为 `{"code":200,"success":true,"msg":"任务控制请求已生效"}`。

上游样例尚无 taskId 或执行记录 ID，也只给出了 running 状态。稳定执行标识、暂停/恢复的实际字符串、同周期重新检查和重复请求语义需联调确认，不能把 proposed paused/resumed 当成已支持协议。建议采用协作式暂停：单次已开始的模型请求允许收束，在文件/页检查点停止后续工作，恢复后继续原范围并校验 fence。

需区分两个 URL：上游调用本服务的 check_url 与服务发送结果的 resultCallbackUrl。前者暂停/恢复时沿用原执行记录 URL；后者随本轮任务保存快照，不能在投递途中被新请求静默替换。

### 3.2 结果回调

正式回调必须映射为：

```json
{
  "mineId": "610203006642",
  "itemCode": "GAS_DAILY_001",
  "checkCycle": "2026-10",
  "checkTime": "2026-10-14T14:35:20+08:00",
  "checkResult": "pending",
  "description": "部分资料待核验",
  "detail": "2026-10-14 防冲日报存在预警候选，但 OCR/VLM 未完成交叉核对。",
  "attachments": []
}
```

允许的五类结果为：`pass`、`no_pass`、`pending`、`unverifiable`、`failed`。五类结果都必须有 `description` 和 `detail`。`attachments` 中的 `attachmentContent` 只在附件存在具体问题时填写；没有问题时省略该字段，不发空字符串或 `null`。

当前代码只实现本地 `mock/none` 交付，没有真实 HTTP callback、URL 快照、2xx 确认、重试、幂等和远端回调状态。

上游要求每个检查项每次执行只上报一条完整结果。多个日期的问题集中写进该条 detail，不按日期拆成多次回调。checkCycle 原样回传；attachmentCycle 是附件归属周期，可以不同于检查周期。五种状态都上报，只有 no_pass 计入问题预筛选。下游人工核实不等于本服务内人工改判。

正式交付应同时核对约定 HTTP 成功码和业务成功响应，例如 code=200 且 success=true。HTTP200但success=false不算交付。超时可能是远端已收、应答丢失，所以重试必须复用同一 payload 和幂等标识。样例无业务幂等字段，应协商执行标识或HTTP请求头，不能自行承诺 exactly-once。回调失败只重投，不重新执行检查。

## 四 OceanBase 清单取证方案

### 4.1 连接边界

用户提供的目标环境是 OceanBase：

```text
driver: com.oceanbase.jdbc.Driver
host: 172.26.10.25
port: 2883
database: sy_xian_mine_dev
user: colliery@co
```

密码不得写入本交接文档、代码、日志或测试快照。生产实现使用环境变量或密钥管理系统。当前服务没有 OceanBase 适配器，也没有读取 `mine_remote_disaster_file` 的代码。

### 4.2 查询规则

服务端维护 `itemCode -> file_name_id -> material_group -> role -> cadence` 白名单。客户端只能传矿井、周期和检查项，不能任意传入 `file_name_id` 或任意 URL。

查询逻辑的最小约束：

- `mine_code` 来自可信上游主体映射；
- `report_time_year` 与检查周期计算一致；
- 月份或日期范围必须显式记录；
- `file_name_id` 由服务端检查项目录决定；
- 返回原始行、查询参数、查询时间和查询版本；
- 按 `attachment_id` 去重；
- 同名不同 ID 保留；
- 同一资料类别的不同内容哈希版本保留并标记冲突；
- 原始 URL、附件名、报表日期、清单行和字节哈希进入本轮快照。

清单索引只说明资料存在或被登记，不能直接证明预警、撤人、卸压或复工行为。

### 4.3 服务端资料目录

以下来自用户提供的SQL，属于接入配置初稿，需分别映射两项规则。SQL中的若干ID之间缺少逗号，开发应生成参数化IN占位符列表，不原样复制执行。

| file_name_id | 清单类型 | 主要用途 |
| --- | --- | --- |
| `rock_ledger_1k8hqc4_rpt_wh5nwe` | 防冲工作日报表 | 预警线索和当日处置 |
| `base_base_ledger_1falhnq_x_smfptn` | 调度日报 | 调度安排与生产记录 |
| `rock_ledger_1k8hqc4_rec_5sdl97` | 日分析记录 | 危险研判与负责人确认 |
| `rock_ledger_1k8hqc4_x_1t0mx6a` | 生产组织通知单 | 处置或复工指令 |
| `rock_ledger_1k8hqc4_rec_axez36` | 防冲钻孔验收记录 | 施工参数及验收 |
| `rock_rpt_181s4w6_x_jqiyc1` | 预警临界指标 | 生效阈值和比较方向 |
| `rock_x_iverp1_x_1t8xner` | 防冲设计 | 适用义务与解危要求 |
| `rock_rpt_181s4w6_mining_rpt_1vaqrs5` | 工作面危险性评价 | 工作面与危险区域 |
| `rock_rpt_181s4w6_rpt_1f3n8vh` | 煤层危险性评价 | 矿井和煤层适用范围 |

需确认OceanBase租户兼容模式和可用驱动。Java driver-class-name不能直接导入Python；正式Python驱动应在实际环境验证后选择，不为此先新增旁路服务。检查服务只读附件表，结果统一回调上游，不直接写上游问题表。

### 4.4 日期和完整性

checkCycle、附件登记日期、报告业务日期、事件时间、上传时间、设计有效期分别保存。年、半年、周、月是上传频率，不自动等于适用有效期。真实清单中的2026/01/00、2026/07/00是周期标签，不是事件日期，不能补成某日。

月度任务先取月内日报和原始记录，再按生效范围取设计与阈值，不能用月度过滤误删年度资料。跨月未闭合事件可扩展证据查询边界，但保持原checkCycle不变并说明扩展原因。周验收以正文施工日期为准，不能用登记日期证明后续几天已卸压。查询失败、没有返回行、下载失败、解析失败、预算跳过、已完整检索须分开记录。范围完整性只针对已声明资料，不增加催报或上传管理流程。

## 五 防冲日报线索和两项检查逻辑

### 5.1 线索字段

防冲工作日报是第一项检查的线索入口。每条线索应尽量形成：

```text
clue_id
date
workface 或巷道
event_type: warning / dynamic
event_key（有则保留）
continuation_key（有则保留）
metric、value、unit、level、method
certainty: confirmed / candidate
evidence_refs
source_file_id
page/row/cell locator
quote
```

明确“未发生预警”“未出现动力现象”的句子不能生成预警线索。泛化的“预警指标”“监测制度”“应当预警”不能生成实际事件。没有日期或工作面时保留候选并记录绑定缺口，不能跨文件强行关联。

原生文本或已完成字段级双候选核对的 OCR/VLM 片段才可进入 `confirmed`。未经核对的模型片段只能进入 `candidate`，不能直接产生 `no_pass`。

confirmed必须明确表示哪一层已确认：转录一致、字段值确认、事件发生确认不能共用一个含糊标记。两个模型转录相同不等于确实发生预警。曲线坐标、图例强中弱、危险等级和当前系数，须绑定监测对象、时间和有效阈值或明确事件记录，不能仅凭“危险等级”确认触发。优先定位日报结论区、报警表和措施区，复杂曲线另做坐标/图例绑定。

### 5.2 安排状态

撤人和卸压安排从调度日报、日分析、生产组织通知、钻孔验收等实际记录中提取：

```text
arranged              明确已安排/已执行/已落实
explicit_not_arranged 明确写出未安排/未执行/未落实
planned_only          只有计划、应当、拟执行、准备执行
evidence_not_found    本次资料范围内没有找到安排证据
unknown               有冲突、缺日期、缺工作面、OCR 不可读或无法判断
```

要求或批准设计属于 `requirement`，不能证明现场已经安排或执行。`evidence_not_found` 不能改写成 `explicit_not_arranged`。

### 5.3 日问题和连续事件

每条日报线索先输出逐日问题：

```text
2026-10-14 发现预警线索，工作面为 403105 工作面。
撤人：本次资料未找到撤人安排证据。
卸压：本次资料未找到卸压安排证据。
```

只有以下条件同时成立才允许合并连续周期：

- 同矿井；
- 同工作面或巷道；
- 同事件类型；
- 有显式事件编号或连续事件标识；
- 日期连续；
- 事件没有明确解除/结束标记。

相邻日期本身不能证明是同一事件。无编号的相邻日报应保留多个问题周期。闭合事件不能和后续日期自动合并。

### 5.4 版本冲突

同一矿井、同一报表日期、同一资料类别存在多个附件版本时：

- 保留全部附件 ID、文件名、哈希和证据；
- 内容哈希相同的重复附件不升级为版本冲突；
- 内容哈希不同的版本标记 `version_conflict`；
- 相关问题最高只能是 `unverifiable`，不能确定 `no_pass`；
- 只有上游提供明确作废/生效关系后，才允许选择权威版本。

### 5.5 处置义务和核验范围

线索只负责召回可能事件。确定判断需核对适用区域、监测方法、实际值、单位、比较方向、批准阈值和事件时间。预警线索不能单独证明法定触发条件。卸压义务还依赖危险研判和适用设计，不能对每次预警一律要求卸压；触发已确认但没有研判时，输出责任/时序缺口。

安排、下令、实施、完成、验收分别保存。调度已安排撤人不等于全部人员已撤离；设计规定孔深不等于已施工。继续采掘反证限于同区域危险时段；设备运行、检修、同日有产量不能自动证明带险采掘。日期粒度不能证明具体动作先后，不自行增加五分钟等期限。

效果检验检查可由明确解危施工完成记录建立检验事件，继续找效果原始测值、必要测点、阈值、确认、解除和实际复工，不能因没有统一预警编号而直接丢弃施工事实。等于阈值不满足严格小于的恢复条件，具体问题和重大隐患的比较符须按法律前件分别核实，不把 > 和 >= 混用。

## 六 结果状态和上报语义

| 状态 | 含义 | 允许进入 no_pass 问题预筛选 |
| --- | --- | --- |
| `pass` | 必要事实、关联和规则命题完整闭合，未发现问题 | 否 |
| `no_pass` | 已确认线索和明确反证形成完整问题证据链 | 是 |
| `pending` | 资料范围或事件关联不完整，但存在可继续核验的事实 | 否 |
| `unverifiable` | OCR/VLM、文件质量或关键内容无法可靠读取 | 否 |
| `failed` | 技术执行失败，未形成可用业务快照 | 否 |

对“撤人”和“卸压”的结果描述必须分别写出状态。示例：

```text
2026-10-14 发现预警线索；明确未安排撤人；本次资料未找到卸压安排证据。
```

“明确未安排”是必要反证之一，但不能单凭此短语或 confirmed 标记生成确定问题。还必须证明该动作确有义务、适用对象和时段一致、事件关联成立且无排除冲突。候选线索、版本冲突和关键不可读内容不能提升为确定性问题。其他独立完整证据已证明违规时，保留 no_pass 并列其他缺口，不被无关扫描失败掩盖。

## 七 顶尖专家会补充的原始素材

当前附件清单之外，专家会要求：

- 微震、地音、应力、钻屑、电磁等原始时序数据；
- 报警计算日志和阈值配置版本；
- 防冲日报原始扫描页和原始表格；
- 调度指令、停送电记录、人员定位、入井考勤和设备运行；
- 撤人名单、路线、人数和避灾地点；
- 卸压钻孔轨迹、孔深、角度、间距、爆破记录和验收原件；
- 效果检验原始测点、测值、单位、采样时间和阈值；
- 技术负责人签署、调度解除和实际复工指令；
- 评价报告、防冲设计、临界指标和有效期版本；
- 上游历史人工核查、整改闭环和作废附件审计。

## 八 结果正确性的独立证明

`completed`、模型置信度和服务自测不能证明结论正确。需要：

- 当前有效法规和矿井批准设计；
- 调度、停送电、人员定位、考勤、生产等独立系统交叉证据；
- 两名防冲专家独立标注，分歧由第三人仲裁；
- 原件字节哈希、页/行/单元格定位、规则版本和解析版本；
- 可重放的查询、输入快照、证据快照和结果版本；
- 按矿井、格式、事件类型和缺口分层的独立 gold/盲测集；
- 误报、漏报、事实召回、证据定位率、不可读率、版本冲突识别率和连续事件归并准确率。

## 九 当前系统审计

### P0 接入缺口

1. 没有用户提供的 `mineId/status/items/resultCallbackUrl` 任务接收接口。
2. 没有 OceanBase 查询适配器和检查项资料白名单。
3. 没有真实 HTTP 结果回调、回调 URL 快照、2xx 确认、重试和远端幂等。
4. 没有暂停/恢复任务状态机。
5. 本地 `sourceSystem/mineCode/environment` 是客户端自报，不能作为生产鉴权。

### P1 业务和质量缺口

1. 当前规则登记哈希与用户当前规则文件哈希不一致。
2. 防冲日报 OCR/VLM 不稳定时只能得到候选，真实线索仍不能稳定进入确定事实。
3. 真实样本中已经有原文和候选，但事件链经常为 0，不能输出最终问题。
4. 10 月 15 日双版本只完成了保留和冲突识别，尚未接入源端作废/权威版本。
5. 当前测试主要是工程和合成样本，不是生产准确率验收。

### 已完成能力

1. 原生 PDF、DOCX、XLSX 解析和文件字节哈希。
2. 混合 PDF 原生文字保留，实际图片区域 OCR/VLM 候选和页面/provider 进度。
3. XLSX 标题行、真实表头、行/单元格来源定位和日级生产事实抽取。
4. 事实证据链、候选关联、日报线索、安排状态和连续周期纯函数。
5. 同日版本冲突保留、调试页、进度查询和本地 SQLite 账本。
6. 前次执行记录为全量回归 256 项通过；本次文档编制未重跑全部测试。该数量不代表真实业务准确率。

### 9.1 本次审阅确认的具体缺陷

本次没有改规则代码。以下短句可复现当前提取边界，开发应加入回归测试：

| 输入 | 当前行为 | 调整要求 |
| --- | --- | --- |
| `403105工作面危险等级弱` | 可生成 confirmed warning 线索 | 等级标签只能是候选，须有事件/阈值前件 |
| `403105工作面未记录动力现象` | 可生成 confirmed dynamic 线索 | 否定/未记录不能成为发生事实 |
| `403105工作面未找到撤人记录` | 可生成 confirmed not_arranged | 应为 evidence_not_found 或 unknown |
| `service.py` confirmed_daily 后处理 | 可在规则评估后追加 no_pass | 不能绕过适用义务、阈值、研判和时序 |

还需验：同页多工作面、同页多日期、缺矿井/缺工作面、同源候选重复、版本冲突、同日两个独立事件、解除后再次预警。当前 `daily_*` 后处理应和 rules.py 统一成一条命题链，不能形成第二个覆盖式判定入口。

### 9.2 真实耗时和失败恢复

真实轮次已经证明模型请求耗时和网络波动会占主要时间。需要把稳定原生解析、成功模型候选、临时 provider 失败分开缓存；失败页可单独重试，不能把临时 URLError 固化成最终快照。模型、提示、token预算、质量门和解析版本应参与识别缓存身份。页预算跳过必须出现在结果和回调详情中。关闭 OCR 后完成的运行只能作流程对照，不能作真实扫描验收。

当前本地调试页可以显示文件、页、provider 和进度事件；生产接入还需输出 task/run/item 三层状态、查询耗时、下载耗时、解析耗时、模型耗时和回调耗时。

## 十 开发调整方案

下列 M0 到 M6 沿用现有阶段名，不重新建设大平台。“冻结”只表示为本轮形成一致契约并登记版本，不增加审批，也不阻止后续修订。建议实际顺序为：先修判定语义和规则版本，再贯通一条真实上游任务闭环，再扩展样本和恢复能力。

### M0 接口和身份契约

- 新增上游任务下发适配器；
- 保存任务原文、任务 ID、状态、周期、检查项和回调 URL；
- 任务/检查项幂等；
- 明确暂停/恢复语义；
- 建立可信主体到矿井、任务、运行和证据的授权链；
- 冻结规则版本和检查项资料目录。

### M1 OceanBase 和附件取证

- 只读 OceanBase repository；
- 参数化查询；
- `itemCode -> file_name_id` 白名单；
- 查询结果快照；
- 附件 ID 去重和版本保留；
- 流式下载、文件魔数、SHA256 和重定向约束；
- 记录缺失、证书失败和附件损坏。

### M2 事实层和识别质量

- 防冲日报线索字段标准化；
- 普通 OCR、版面 OCR、VLM 候选分层；
- 字段级交叉核对；
- 同页多工作面分段；
- 表格、签名和图片覆盖清单；
- 候选、事实、缺口和不可读状态分开保存。

### M3 规则和周期问题

- 预警线索状态机；
- 撤人、停工、断电、报告、卸压的关联；
- 明确未安排与资料未找到证据；
- 多日期问题；
- 连续事件归并；
- 版本冲突；
- 无事件关联不能直接判定不合规。

### M4 结果回传

- 对外 callback DTO；
- `mineId`、`itemCode`、`checkCycle`、`checkTime` 映射；
- 附件问题内容和周期映射；
- durable callback outbox；
- 2xx 确认、重试、幂等和审计；
- 回调失败不重新执行检查。

### M5/M6 真实验收

- 真实矿井 gold；
- 专家双标和仲裁；
- 开发集/盲测集隔离；
- 分层准确率和召回率；
- 证据定位率；
- 不可读率；
- 版本冲突识别率；
- 连续事件归并准确率；
- 重跑一致性和正式上游联调。

## 十一 开发验收清单

- [ ] 上游任务下发响应成功只表示控制动作生效，不表示结果已产生。
- [ ] 同一任务重复下发不会创建重复运行。
- [ ] 每个检查项独立执行、独立回调。
- [ ] OceanBase 查询只使用服务端检查项白名单。
- [ ] 查询快照和附件字节版本可重放。
- [ ] 10 月 15 日双版本被保留并显示版本冲突。
- [ ] 日报中的候选线索不会直接产生 `no_pass`。
- [ ] 明确“明确未安排”和“本次资料未找到证据”。
- [ ] 同工作面、显式事件标识、连续日期才归并 episode。
- [ ] 缺日期、缺工作面、OCR 不可读时不会跨文档强拼。
- [ ] 每个问题都能定位到文件、页/行/单元格和原文。
- [ ] `no_pass` 回调只引用确定性证据。
- [ ] `pending/unverifiable/failed` 都有完整 description/detail。
- [ ] 回调失败重试同一 payload，不重新检查。
- [ ] 跨矿、跨运行和跨环境读取被拒绝。
- [ ] 真实专家 gold 和盲测结果达到放行阈值。

验收阈值由业务方与独立专家在固定样本和指标分母后确定，不在本文件臆定准确率。开发阶段先用小型独立gold检查错误，M6再用隔离盲测放行。上游实收验收须分别证明控制响应成功和每项结果回调业务成功，不能用mock_receipts代替。

## 十二 当前真实验证记录

- `1024.xlsx`：胡家河矿业，清单 1173 行，唯一附件 ID 1089，去重重复行 84 行。
- 10/14–10/16 轮次：防冲日报、调度日报、日分析和 10/15 双调度版本；真实模型处理得到候选日报线索，但因 OCR/VLM 不稳定保持 candidate/unverifiable。
- 9/23–9/25 轮次：流程可受理和收束，但关闭模型快速验证没有业务线索结论。
- 10/20–10/22 轮次：部分附件下载失败，未提交检查任务，避免把数据缺失误报为业务未安排。
- `西村矿业all.xlsx`：表内实际矿井为孟村矿业；6 个真实附件完成了流程验证，产生 4 条日级生产事实，但没有确认的危险触发事件链。

以上记录说明工程链路可运行，但真实业务准确性、上游接入和 OceanBase 取数仍需按本交接方案完成。

## 十三 交接文件和来源

- 规则源：`D:\西安院\深度监察\冲击地压\冲击地压规则1023.xlsx`
- 上游回调示例：`D:\西安院\深度监察\冲击地压\检查结果上报接口示例.md`
- 当前服务：`compliance-service`
- 日报线索说明：`compliance-service/docs/daily-clue-flow.md`
- 1024 选样脚本：`compliance-service/scripts/prepare_1024_rounds.py`
- 多轮参数：`compliance-service/.local/validation/1024-multi-round/`
- 真实样本清单来源：`C:\Users\86158\Desktop\1024.xlsx`

## 十三 独立正确性证明

可证明结果正确的证据不包括 completed 状态或模型置信度。应使用当前法规、批准防冲设计、调度/停送电/人员定位/考勤/生产等独立系统、原始附件字节哈希、页行单元格定位、查询和输入快照、专家双标及仲裁、冻结开发 gold 和独立盲测集。统计字段召回率、事件召回率、问题精确率、误报/漏报、证据定位率、不可读率、版本冲突识别率和连续事件归并准确率。

顶尖专家不会把清单目录当作行为事实。他们会从微震、地音、应力、钻屑等原始时序和报警计算日志开始，并把防冲日报、调度记录、停送电、人员定位、撤人名单、卸压参数、效果测值、签署和解除指令交叉核对。每种资料都应标明“可以证明什么”和“不能证明什么”。

## 十四 开发题目改写

**基于上游任务与 OceanBase 取证的冲击地压两项跨文档合规检查服务调整方案**

请沿用当前服务，接收检查项基础表 `check_url` 对应的上游任务报文，按 `mineId`、`checkCycle` 和 `items` 从 OceanBase 查询服务端维护的资料目录，冻结本轮输入范围和附件版本，解析防冲工作日报、调度日报、生产矿长/总工程师日分析、防冲设计、预警临界指标、钻孔验收和评价报告。

对“停工撤人与卸压”，先从防冲工作日报识别带日期、工作面、事件类型和证据定位的预警/动力线索，再以同矿井、同日期、同工作面和明确事件标识为约束，关联调度、日分析、生产组织通知和卸压资料，分别判断撤人和卸压。对每个日期输出明确已安排、明确未安排、仅有计划、本次资料未找到证据、不可读、版本冲突和事件关联不足。相邻日期只有在存在明确连续事件标识时才合并。

每个检查项独立执行并通过上游指定回调地址回传一条完整结果。结果包含矿井、检查项、检查周期、检查时间、五类结论之一、摘要、详细依据和附件引用。服务不得提供人工改判接口；历史结果不可变，修订通过新运行或新结果版本完成。

## 十五 开发交付物

开发完成后必须提交：

- 上游任务接口、状态和幂等契约；
- OceanBase 只读查询实现和检查项资料目录；
- 查询快照、附件版本和字节哈希数据结构；
- PDF/XLSX/DOCX 事实和证据结构；
- 日报线索、安排、问题、episode 和版本冲突实现；
- 内部结果到回调 DTO 的映射；
- callback outbox、重试、2xx 确认和幂等实现；
- 任务、文件、页、provider 和回调进度；
- 迁移脚本、配置说明和回归测试；
- 真实案例的输入快照、结果 JSON、专家依据和已知限制；
- 真实 gold/盲测评测报告。

服务交付文档必须明确哪些已联调、哪些只经过本地模拟、哪些仍等待上游或 OceanBase 环境。不能用模拟结果、静态页面或“completed”状态代替真实生产链路证明。

## 十六 代码定位

| 文件 | 当前职责 | 开发重点 |
| --- | --- | --- |
| `compliance_service/api.py` | 本地 manual/upstream/debug 入口 | 增加真实任务适配器和回调状态，不破坏本地调试入口 |
| `compliance_service/contracts.py` | 内部请求、运行和快照约束 | 分离上游 DTO 与内部 RunSpec，增加任务/项幂等 |
| `compliance_service/sources.py` | manifest 文件取证、解析缓存、事实快照 | 接 OceanBase 查询快照和来源元数据 |
| `compliance_service/fact_extraction.py` | 结构化事实和证据引用 | 增加原始监测、线索、签署和安排语义，避免泛文本误提 |
| `compliance_service/daily_clues.py` | 防冲日报线索和安排候选 | 接可靠候选、质量状态和日期/工作面上下文 |
| `compliance_service/daily_episodes.py` | 多日期问题和连续事件归并 | 接版本、作废、同源重复和冲突 |
| `compliance_service/rules.py` | 两项时序与阈值规则 | 统一日报问题和规则命题，不能形成第二个绕过入口 |
| `compliance_service/ledger.py` | SQLite 账本、租约、outbox、进度 | 生产存储评估、任务控制和真实 callback outbox |
| `compliance_service/ocr_documents.py` | 页/图片识别和候选质量门 | 重试、预算、provider稳定性、候选与确认分层 |
| `scripts/prepare_1024_rounds.py` | 1024.xlsx 多轮选样 | 维护真实案例选择和来源元数据 |
| `scripts/run_1024_trials.py` | 真实附件下载/校验/重跑 | 仅用于本地验收，不替代生产 OceanBase 接入 |

交接时不得复制历史聊天中的数据库密码或模型密钥。连接参数和凭据分别由部署环境的密钥管理提供。
'''

# Keep the written source and generated Word document synchronized.
MD = (MD.replace("## 十三 独立正确性证明", "## 十四 方法和正确性验证")
        .replace("## 十四 开发题目改写", "## 十五 开发题目改写")
        .replace("## 十五 开发交付物", "## 十六 开发交付物")
        .replace("## 十六 代码定位", "## 十七 代码定位"))
MD = MD.replace("只有“明确未安排”且线索已确认时", "只有全部规则前件成立且存在明确反证时")
MD = MD.replace('"mineId": "610203006642"', '"mineId": "610401015959"')
MD = MD.replace('"itemCode": "GAS_DAILY_001"', '"itemCode": "stop_work_unload"')
MD = MD.replace('"itemName": "瓦斯日报检查"', '"itemName": "冲击危险停工撤人与卸压"')
MD = MD.replace("## 十三 交接文件和来源", "## 十三 交接文件和来源")
MD = MD.replace("当前代码只实现本地 `mock/none` 交付，没有真实 HTTP callback、URL 快照、2xx 确认、重试、幂等和远端回调状态。", r'''当前代码只实现本地 mock/none 交付。本地 outbox 已有重试和租约机制，但没有真实 HTTP callback 客户端、URL快照和远端业务成功确认。应复用现有机制，而非另建一套交付队列。

对外字段映射如下，内部快照和候选不要直接整体塞给上游：

| 上游字段 | 内部来源 | 约束 |
| --- | --- | --- |
| mineId | 可信任务矿井身份 | 字符串保存 不丢前导零 |
| itemCode | 受理时检查项编码 | 保持与下发一致 |
| checkCycle | 原始任务周期 | 不改成附件周期 |
| checkTime | 本轮结果定版时间 | 带时区 不随投递重试改变 |
| checkResult | 内部五类结果 | 只发送约定枚举 |
| description | 全项摘要 | 包含对象 周期 关键结论 |
| detail | 全部问题与缺口 | 多日期问题逐项列举 |

附件对象示例是字段说明，以下不是实际矿井问题：

```json
{
  "attachmentCycle": "2026-10",
  "attachmentName": "防冲工作日报.pdf",
  "attachmentUrl": "https://files.example.com/evidence.pdf",
  "attachmentContent": "该附件第2页包含触发线索，具体问题见本轮证据链"
}
```

原始附件URL和证据定位应可用于上游查看，不能用本机绝对路径、file_id或内容寻址目录冒充attachmentUrl。无该附件自身具体问题时省略attachmentContent。暂时网络失败和HTTP5xx可退避重试；鉴权/报文错误及业务失败应进入可见失败/等待修复状态，不无限重投，也不因此改写业务结果。''')
MD = MD.replace("预警或动力线索 -> 研判 -> 停工、断电、撤人、报告 -> 卸压措施和验收 -> 继续采掘反证",
                "触发确认 -> 立即停工、断电、撤人、报告 -> 研判 -> 按持续危险核卸压/验收 -> 采掘反证")
MD = MD.replace("- 新增上游任务下发适配器；", "- 在既定check_url路由或网关映射下新增任务适配器，不要求上游更改地址；")
MD = MD.replace("- 任务/检查项幂等；", "- 先确认执行标识，再实现任务/检查项幂等；上游无ID时协商组合键或幂等请求头；")
MD = MD.replace("6 个真实附件完成了流程验证，产生 4 条日级生产事实", "6 个真实附件实际开启模型运行产生 434 个原生/候选证据块和 4 条日级生产事实")
MD = MD.replace("真实模型处理得到候选日报线索，但因 OCR/VLM 不稳定保持 candidate/unverifiable。",
    "真实模型运行 b03bfbe1d19543889866ac670917ddc8 耗时约 537.7 秒，历史结果保存了 3 个候选片段和 3 个周期。它们存在同源重复身份问题，后续去重代码内存重算为 1 条，但历史库未重写；不能把候选片段当成 3 个独立预警事件。")
MD += r'''

## 十八 开发任务和验收用例

| 调整任务 | 最小交付 | 必须取得的验证证据 |
| --- | --- | --- |
| T01 统一规则语义 | F2新revision 否定与缺证决策表 单一判断入口 | 危险等级弱 未记录动力现象 未找到撤人记录三个反例不误判 |
| T02 上游任务接入 | check_url路由 任务项身份 控制状态 幂等 | 上游原样报文受理 重复控制可解释 暂停恢复定位原执行 |
| T03 OB资料查询 | 只读repository 两项清单ID目录 周期有效期策略 | 查询行数和附件ID与上游核对 跨矿请求拒绝 |
| T04 文档事实层 | 输入快照 表头曲线对象绑定 页面覆盖 单元重试 | 日期 工作面 数值 单位和原文逐条可追溯 |
| T05 两项事件核验 | 触发前件 动作时序 多日问题 版本和结束条件 | 候选不越级 缺证不当明确反证 跨日不误借证据 |
| T06 回调闭环 | DTO 定版outbox URL快照 HTTP和业务成功确认 | 五类结果上游实收 超时重投同payload 多问题仍只一条 |
| T07 真实验收 | 小型真实gold 专家复核 M6盲测 | 字段 事件 问题和证据定位指标分母可解释 |

必测场景应覆盖：小于/等于/大于阈值；不同方法、单位和比较方向；只有图例没有事件；仅计划没有实际执行；明确否定与查无记录；同日动作早于触发；仅日期精度；同页多面；跨日显式关联；解除后再次预警；同字节重复附件；相反版本；缺少某一天资料；OB/附件/OCR/视觉失败；恢复仅补失败单元；HTTP200但业务失败；回调已收但应答丢失；跨矿、跨运行证据引用。

用变形用例验证逻辑：删掉必要证据不能提升结论；换成其他矿井或工作面不能闭合；调整文件顺序不改变语义；重复字节不能制造新事件；等价单位转换不改变合法比较；把“未安排”改成“未找到记录”应降为缺证。

## 十九 建议采用的具体工具和最小实现路径

建议使用 OceanBase 兼容的只读数据库客户端及参数化SQL；附件采用流式HTTP、证书校验和SHA256；PDF使用pypdf和Poppler，Excel使用openpyxl，Word使用python-docx。保存文字层、页图、表格行列、合并单元格、原始公式缓存和坐标。模型只负责生成候选，不直接生成业务结论。

事实表、证据表、文件版本表和事件关系可继续采用关系型结构，以时间区间join和显式状态机实现，不需要先建图数据库或通用规则DSL。优先复用现有运行账本、fence、outbox和进度机制；是否更换生产账本由实际部署与并发验证决定，不为扩大架构先引入中间件。必要时再把当前结构化进度接入OpenTelemetry和Prometheus。

专家会先建立资料能证明什么的矩阵，再以原始数据重建时间线。例如，调度指令可证明安排，人员定位和人数名单可证明撤离实施，设计证明义务，验收证明特定施工，检验原始数值证明某时某点结果。签名识别不能仅凭姓名文字认定签署真实性。

正式M6前可以先做少量开发gold，暴露错误而不阻塞开发；最终采用独立专家双标、争议仲裁和按矿井/格式/事件/缺口分层的盲测。不要把pending/unverifiable直接算作正确业务判断，不要把同一case两项运行或同源OCR片段当作独立样本扩大分母。

## 二十 联调前待确认事项

1. 两项正式itemCode及check_url，内部stop_work_unload和relief_effects是否直接作为正式编码；瓦斯日报示例编码不用于这两项。
2. 上游执行标识、status允许值、重复running、终态后再次下发和暂停/恢复定位方式。建议活动运行合并，后续修订创建新运行；最终以双方契约为准。
3. OB租户兼容模式、只读账号权限、实际DDL、行更新时间及附件替换/作废字段；凭据单独交付。
4. 资料有效期、跨月事件和较早周验收的适用策略；没有人员轨迹、停送电或效果原始记录时的可核验边界。
5. 回调身份认证、正式域名/端口、业务成功判据和幂等机制；示例域名不直接用于部署。
6. 原表F2新版本与官方法规的一致性，等于阈值时不同分支的处理，以及范围完整但没有触发事件时pass的表述。

这些事项应形成简明契约和版本记录，不新增服务内人工审批或人工改判页面。本次未完成官方法规逐条法律审定，也未验证真实OB、正式回调和生产网络权限。

## 二十一 本次验证记录的适用范围

三轮历史案例不能合称“3轮真实业务验收通过”：10月中旬开启模型但仍为候选；9月轮次是关闭模型的流程对照；最新轮次附件不全，没有完成有效检查。正式调整后至少以完整资料再验证正例、明确反例和未知例，并保存各层输入和原始结果。

下载失败后不得将成功子集按顺序zip回原附件列表，否则会把A附件字节绑定为B附件。历史选样脚本曾出现该风险，后续可重复脚本改为attachment_id逐项核对。复用历史local-request前必须核对ID、原始文件名、字节哈希和报告周期；魔数正确不等于附件身份正确。

'''


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=11, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    for name, size, bold, color in (
        ("Normal", 11, False, "000000"),
        ("Title", 20, True, "000000"),
        ("Heading 1", 15, True, "000000"),
        ("Heading 2", 11.5, True, "000000"),
        ("Heading 3", 10.5, True, "000000"),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
    doc.styles["Normal"].paragraph_format.space_after = Pt(3)
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)


def add_paragraph(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    # Keep inline code readable without creating a code-heavy visual style.
    parts = re.split(r"(`[^`]+`|\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=9.2, color="374151")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)
    return paragraph


def add_table(doc, lines):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and not all(re.fullmatch(r"[-: ]+", cell) for cell in cells):
            rows.append(cells)
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    available = 6.94
    if columns == 3:
        ratios = (0.44, 0.23, 0.33) if rows[0][0] in {"file_name_id", "文件"} else (0.23, 0.33, 0.44)
    else:
        ratios = tuple(1 / columns for _ in range(columns))
    for column, ratio in zip(table.columns, ratios):
        column.width = Inches(available * ratio)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for index, cell in enumerate(table.rows[0].cells):
        value = rows[0][index] if index < len(rows[0]) else ""
        cell.text = value.replace("`", "")
        set_cell_shading(cell, "244764")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.3, bold=True, color="FFFFFF")
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for row_index, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        for index, cell in enumerate(cells):
            cell.text = values[index].replace("`", "") if index < len(values) else ""
            set_cell_margins(cell)
            set_cell_shading(cell, "F2F5F7" if row_index % 2 == 0 else "FFFFFF")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9.3)
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell, ratio in zip(row.cells, ratios):
            cell.width = Inches(available * ratio)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.08
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_docx():
    doc = Document()
    style_doc(doc)
    lines = MD.splitlines()
    index = 0
    first_heading = True
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            run = paragraph.add_run(line[2:].strip())
            set_run_font(run, size=20, bold=True)
            if first_heading:
                paragraph.paragraph_format.space_after = Pt(16)
                first_heading = False
            index += 1
            continue
        if line.startswith("## "):
            if line.startswith("## 二 用户确认"):
                doc.add_page_break()
            doc.add_heading(line[3:].strip(), level=1)
            index += 1
            continue
        if line.startswith("### "):
            doc.add_heading(re.sub(r"(?<=\d)\.(?=\d)", " ", line[4:].strip()), level=2)
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run("\n".join(code_lines))
            set_run_font(run, size=9.2, color="374151")
            continue
        if line.startswith("| "):
            table_lines = []
            while index < len(lines) and lines[index].startswith("| "):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("- "):
            paragraph = add_paragraph(doc, line[2:].strip(), style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.18)
            index += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            add_paragraph(doc, line)
            index += 1
            continue
        add_paragraph(doc, line)
        index += 1
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    run = footer.add_run("冲击地压跨文档合规检查服务开发交接与验收方案 | 2026-10-25")
    set_run_font(run, size=8, color="6B7280")
    footer.add_run("  第 ")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    footer.add_run(" 页")
    for border in list(doc.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.core_properties.title = "冲击地压跨文档合规检查服务开发交接与验收方案"
    doc.core_properties.subject = "上游任务、OceanBase取证、跨文档规则和结果回传"
    doc.core_properties.author = "技术交接"
    doc.save(DOCX_PATH)


MD_PATH.write_text(MD, encoding="utf-8")
build_docx()
print(MD_PATH)
print(DOCX_PATH)
