"""Reproducible binary input samples for local acceptance, never real mine data."""

from copy import deepcopy
import json
from pathlib import Path


def _record(minute=0, **values):
    return {"工作面": "W1", "时间": f"2026-10-23T08:{minute:02d}:00",
            "关联事件时间": "2026-10-23T08:00:00", **values}


def _action(name, minute, **values):
    return _record(minute, **{"动作": name, "状态": "已完成", "执行人角色": "防冲工", "目标范围": "W1", **values})


def _write_docx(path, records):
    from docx import Document
    document = Document()
    for record in records:
        document.add_paragraph("；".join(f"{key}:{value}" for key, value in record.items()))
    document.save(path)


def _write_xlsx(path, records):
    from openpyxl import Workbook
    keys = list(dict.fromkeys(key for record in records for key in record))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "实际记录"
    sheet.append(keys)
    for record in records:
        sheet.append([record.get(key, "") for key in keys])
    workbook.save(path)


def create_validation_samples(output_dir):
    """Write complete, missing-evidence and proven-violation file sets."""
    output = Path(output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    point = {"method": "应力监测法", "metric": "煤体应力", "point": "P1"}
    parameter = {"name": "孔深", "operator": ">=", "value": "20", "unit": "m"}
    requirement = {"工作面": "W1", "时间": "2026-09-23", "有效起始": "2026-09-23", "有效截止": "2026-12-31",
                   "批准": "是", "适用范围": "W1"}
    policy = [
        {**requirement, "记录类型": "危险性评价", "冲击危险": "是", "危险等级": "中等"},
        {**requirement, "记录类型": "防冲设计", "技术负责人": "张三", "解危方法全集": '["钻孔卸压"]',
         "解危参数要求": json.dumps([parameter], ensure_ascii=False), "解危区域": "W1危险区", "撤离区域": "避灾硐室A",
         "解危限员": "6", "无关设备停运要求": "是"},
        {**requirement, "方法": point["method"], "指标": point["metric"], "测点": point["point"], "实测值": "10", "单位": "MPa"},
        {**requirement, "指标全集": json.dumps([point], ensure_ascii=False)},
    ]
    groups = {
        "01-warning.docx": [_record(**{"记录类型": "动力现象"})],
        "02-dispatch.xlsx": [
            _action("停工", 1, **{"立即": "是"}), _action("断电", 1, **{"立即": "是"}),
            _action("撤人", 1, **{"立即": "是", "路线符合": "是", "撤离区域": "避灾硐室A", "应撤人数": "18", "实撤人数": "18"}),
            _action("报告", 1, **{"立即": "是"}),
        ],
        "03-analysis.docx": [_action("危险研判", 2, **{"危险研判": "有危险", "签字人": "李四", "签名存在": "是",
                                                    "已告知": "是", "告知时间": "2026-10-23T08:02:30"})],
        "04-relief.xlsx": [_action("卸压", 3, **{"方法": "钻孔卸压", "施工参数": '[{"name":"孔深","value":"22","unit":"m"}]',
                                               "解危区域": "W1危险区", "验收通过": "是", "验收人": "王五",
                                               "验收时间": "2026-10-23T08:03:30", "解危人数": "4", "无关设备已停": "是"})],
        "05-inspection.xlsx": [_action("检验", 4), _record(5, **{"阶段": "效果检验", "方法": point["method"],
                                                                  "指标": point["metric"], "测点": point["point"], "实测值": "8", "单位": "MPa"})],
        "06-release.docx": [_action("技术确认", 6, **{"执行人角色": "技术负责人", "签字人": "张三", "签名存在": "是", "签署时间": "2026-10-23T08:06:00"}),
                              _action("调度放行", 7, **{"执行人角色": "调度室"}), _action("复工", 8)],
        "07-design.docx": policy,
    }
    cases = []
    for scenario in ("complete", "missing-signature", "bad-effect", "bad-relief"):
        folder = output / scenario
        folder.mkdir(parents=True, exist_ok=True)
        records = deepcopy(groups)
        if scenario == "missing-signature":
            records["06-release.docx"][0].pop("签名存在")
        if scenario == "bad-effect":
            records["05-inspection.xlsx"][1]["实测值"] = "10"
        if scenario == "bad-relief":
            records["04-relief.xlsx"][0]["施工参数"] = '[{"name":"孔深","value":"12","unit":"m"}]'
        entries = []
        for name, rows in records.items():
            (_write_xlsx if name.endswith(".xlsx") else _write_docx)(folder / name, rows)
            entries.append({"file_id": Path(name).stem, "path": name,
                            "role": "requirement" if name == "07-design.docx" else "actual_record"})
        manifest = {"source_system": "acceptance-samples", "mine_code": "SAMPLE-MINE", "environment": "test",
                    "period": {"start": "2026-10-01", "end": "2026-10-31"}, "source_complete": True, "files": entries}
        path = folder / "manifest.json"
        path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        for item in ("stop_work_unload", "relief_effects"):
            expected = "pass"
            if scenario == "missing-signature" and item == "relief_effects":
                expected = "pending"
            elif scenario == "bad-effect" and item == "relief_effects":
                expected = "no_pass"
            elif scenario == "bad-relief" and item == "stop_work_unload":
                expected = "no_pass"
            cases.append({"case_id": scenario + "-" + item, "manifest": str(path), "itemCode": item,
                          "checkCycle": "2026-10", "expected_status": expected, "origin": "synthetic_document"})
    (output / "cases.json").write_text(json.dumps(cases, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"sample_kind": "synthetic acceptance documents, not real mine records", "case_file": str(output / "cases.json"),
            "complete_manifest": str(output / "complete/manifest.json"), "case_count": len(cases)}
