"""Real DOCX/XLSX bytes with independently declared document-case expectations."""

from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from compliance_service.contracts import DOCUMENT_DEFINITION_VERSION
from compliance_service.evaluation import summarize_cases
from compliance_service.ledger import Ledger
from compliance_service.service import InspectionService
from compliance_service.sources import LocalDocumentSource, _io_path


HEADERS = ["事件编号", "工作面", "时间", "记录类型", "动作", "状态", "立即", "执行人角色", "目标范围",
           "路线符合", "阶段", "方法", "指标", "测点", "实测值", "单位", "有效起始", "有效截止", "批复", "危险研判", "指标全集"]


def row(minute=0, **values):
    return {"事件编号": "E1", "工作面": "W1", "时间": f"2026-10-23T08:{minute:02d}:00", **values}


def action(name, minute, **values):
    return row(minute, **{"动作": name, "状态": "已完成", "执行人角色": "防冲工", "目标范围": "W1", **values})


def effect_rows(value="8"):
    return [row(**{"记录类型": "动力现象", "危险研判": "有危险"}), action("卸压", 1), action("检验", 2),
            row(3, **{"阶段": "效果检验", "方法": "应力监测法", "指标": "煤体应力", "测点": "P1", "实测值": value, "单位": "MPa"}),
            action("技术确认", 4, **{"执行人角色": "技术负责人"}),
            action("调度放行", 5, **{"执行人角色": "调度室"}), action("复工", 6)]


def stop_rows():
    return [row(**{"记录类型": "动力现象", "危险研判": "有危险"}),
            *(action(name, 1, **{"立即": "是", **({"路线符合": "是"} if name == "撤人" else {})})
              for name in ("停工", "断电", "撤人", "报告")), action("卸压", 5)]


def policy_rows():
    point = {"method": "应力监测法", "metric": "煤体应力", "point": "P1"}
    return [row(**{"方法": "应力监测法", "指标": "煤体应力", "测点": "P1", "实测值": "10",
                   "单位": "MPa", "有效起始": "2026-09-23", "批复": "已批复"}),
            row(**{"指标全集": json.dumps([point], ensure_ascii=False), "批复": "已批复"})]


def save_docx(path, rows):
    document = Document()
    table = document.add_table(rows=1, cols=len(HEADERS))
    for cell, header in zip(table.rows[0].cells, HEADERS):
        cell.text = header
    for values in rows:
        for cell, header in zip(table.add_row().cells, HEADERS):
            cell.text = str(values.get(header, ""))
    document.save(path)


def save_xlsx(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for values in rows:
        sheet.append([values.get(header, "") for header in HEADERS])
    workbook.save(path)


class DocumentPipelineTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.ledger = Ledger(self.root / "state" / "service.sqlite3")
        self.service = InspectionService(self.ledger)

    def manifest(self, actual_rows, *, complete=True):
        # The actual event and later actions deliberately span two source files.
        save_docx(self.root / "event.docx", actual_rows[:1])
        save_xlsx(self.root / "actions.xlsx", actual_rows[1:])
        save_docx(self.root / "policy.docx", policy_rows())
        path = self.root / "manifest.json"
        manifest = {"source_system": "documents", "mine_code": "mine-a", "environment": "test",
                    "period": {"start": "2026-10-01", "end": "2026-10-31"}, "source_complete": complete,
                    "files": [{"file_id": "event", "path": "event.docx", "role": "actual_record"},
                              {"file_id": "actions", "path": "actions.xlsx", "role": "actual_record"},
                              {"file_id": "policy", "path": "policy.docx", "role": "requirement"}]}
        path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        return path

    def submit(self, manifest, item="relief_effects"):
        return self.service.submit({"sourceSystem": "documents", "mineCode": "mine-a", "environment": "test",
                                    "checkCycle": "2026-10", "scope": {"manifest": str(manifest)},
                                    "callbackMode": "none", "items": [{"itemCode": item}]})["items"][0]["run_id"]

    def get(self, run_id):
        return self.ledger.get(run_id, "documents", "mine-a", "test")

    def test_real_docx_xlsx_cross_file_chains_and_evidence_locations(self):
        for item, records in (("stop_work_unload", stop_rows()), ("relief_effects", effect_rows())):
            with self.subTest(item=item):
                run_id = self.submit(self.manifest(records), item)
                result = self.service.run_once()
                self.assertEqual("pending", result["status"], result)
                self.assertEqual("pending", result["implemented_condition_status"])
                state = self.get(run_id)
                self.assertEqual(DOCUMENT_DEFINITION_VERSION, state["spec"]["definition_version"])
                self.assertEqual("local-documents", result["profile"])
                self.assertTrue(all(e["locator"] for e in state["snapshot"]["evidence"]))
                self.assertTrue(all(Path(f["original_path"]).exists() for f in state["snapshot"]["files"]))
                self.assertEqual("none", state["delivery_state"])

    def test_equal_threshold_resumed_is_no_pass_missing_records_pending(self):
        cases = []
        for label, records, expected in (("equal-resumed", effect_rows("10"), "no_pass"),
                                          ("missing-dispatch", [r for r in effect_rows() if r.get("动作") != "调度放行"], "pending")):
            run_id = self.submit(self.manifest(records))
            result = self.service.run_once()
            self.assertEqual(expected, result["status"], result)
            cases.append({"case_id": label, "expected_status": expected, "result": result,
                          "snapshot": self.get(run_id)["snapshot"], "origin": "synthetic_document"})
        report = summarize_cases(cases)
        self.assertEqual(1, report["label_accuracy"])
        self.assertEqual(0, report["unsafe_pass_count"])
        self.assertEqual(0, report["untraceable_case_count"])

    def test_snapshot_copies_and_cache_reuse_do_not_hide_changed_input(self):
        path = self.manifest(effect_rows())
        first_id = self.submit(path)
        self.service.run_once()
        first = self.get(first_id)
        second_id = self.submit(path)
        self.service.run_once()
        second = self.get(second_id)
        self.assertEqual(3, second["snapshot"]["cache"]["reused"])
        self.assertNotEqual(first_id, second_id)
        self.assertTrue({e["evidence_id"] for e in first["snapshot"]["evidence"]}.isdisjoint(
            {e["evidence_id"] for e in second["snapshot"]["evidence"]}))
        # Change only the actual measurement file. Rewriting unchanged DOCX ZIPs
        # may legitimately alter their byte hash through archive timestamps.
        save_xlsx(self.root / "actions.xlsx", effect_rows("11")[1:])
        third_id = self.submit(path)
        self.assertEqual("no_pass", self.service.run_once()["status"])
        third = self.get(third_id)
        self.assertEqual(1, third["snapshot"]["cache"]["parsed"])
        self.assertEqual(first["result"], self.get(first_id)["result"])
        self.assertNotEqual(first["snapshot"]["files"][1]["content_hash"], third["snapshot"]["files"][1]["content_hash"])

    def test_missing_file_or_explicit_partial_scope_cannot_pass(self):
        path = self.manifest(effect_rows(), complete=False)
        self.submit(path)
        self.assertEqual("pending", self.service.run_once()["status"])
        self.manifest(effect_rows())
        (self.root / "actions.xlsx").unlink()
        run_id = self.submit(path)
        result = self.service.run_once()
        self.assertEqual("pending", result["status"], result)
        self.assertEqual(3, len(self.get(run_id)["snapshot"]["files"]))

    def test_manifest_wrong_mine_and_path_escape_cannot_supply_evidence(self):
        path = self.manifest(effect_rows())
        original = json.loads(path.read_text(encoding="utf-8"))
        other = deepcopy(original)
        other["mine_code"] = "another-mine"
        path.write_text(json.dumps(other), encoding="utf-8")
        self.submit(path)
        self.assertEqual("failed", self.service.run_once()["status"])
        original["files"][0]["path"] = "../outside.docx"
        path.write_text(json.dumps(original), encoding="utf-8")
        self.submit(path)
        self.assertEqual("pending", self.service.run_once()["status"])

    def test_corrupted_parse_cache_is_rebuilt_from_immutable_bytes(self):
        path = self.manifest(effect_rows())
        self.submit(path)
        self.assertEqual("pending", self.service.run_once()["status"])
        caches = list((self.root / "state" / "artifacts").glob("*/parsed/*.json"))
        self.assertEqual(3, len(caches))
        corrupted = json.loads(caches[0].read_text(encoding="utf-8"))
        corrupted["payload"]["blocks"] = []
        caches[0].write_text(json.dumps(corrupted), encoding="utf-8")
        run_id = self.submit(path)
        self.assertEqual("pending", self.service.run_once()["status"])
        state = self.get(run_id)
        self.assertEqual(1, state["snapshot"]["cache"]["invalidated"])
        self.assertEqual(1, state["snapshot"]["cache"]["parsed"])

    def test_real_document_workflow_does_not_open_network_sockets(self):
        path = self.manifest(effect_rows())
        with patch("socket.socket", side_effect=AssertionError("offline workflow must not open sockets")):
            self.submit(path)
            self.assertEqual("pending", self.service.run_once()["status"])

    def test_real_document_processing_renews_lease_during_slow_source(self):
        from threading import Event
        clock = [1000.0]
        self.ledger = Ledger(self.root / "renewal.sqlite3", clock=lambda: clock[0])
        self.service = InspectionService(self.ledger)
        path = self.manifest(effect_rows())
        run_id = self.submit(path)
        underlying_source = self.service.source
        renewed = Event()
        renew = self.ledger.renew
        calls = []
        def observed_renew(*args):
            result = renew(*args)
            calls.append(result)
            renewed.set()
            return result
        self.ledger.renew = observed_renew
        def delayed_source(run):
            for now in (1000.6, 1001.2):
                renewed.clear()
                clock[0] = now
                self.assertTrue(renewed.wait(3), "source processing should renew its lease")
            return underlying_source(run)
        self.service.source = delayed_source
        result = self.service.run_once(lease_seconds=1)
        self.assertEqual("pending", result["status"])
        self.assertEqual("pending", result["implemented_condition_status"])
        self.assertGreaterEqual(len(calls), 2)
        self.assertEqual("completed", self.get(run_id)["execution_state"])

    def test_deep_storage_paths_keep_originals_and_parse_cache(self):
        path = self.manifest(effect_rows())
        run_id = self.submit(path)
        # The temporary directory itself also uses an extended path so its
        # normal context-manager cleanup can traverse long Windows filenames.
        with TemporaryDirectory(dir=str(_io_path(self.root))) as temporary:
            storage = Path(temporary) / ("a" * 70) / ("b" * 70)
            self.service.source = LocalDocumentSource(storage)
            result = self.service.run_once()
            self.assertEqual("pending", result["status"], result)
            self.assertEqual("pending", result["implemented_condition_status"])
            self.assertEqual(3, self.get(run_id)["snapshot"]["cache"]["parsed"])

