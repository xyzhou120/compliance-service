"""Real native parsing through task intake, durable control, and HTTP outbox."""
import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from compliance_service.acquisition import AttachmentAcquisition
from compliance_service.api import create_app
from compliance_service.callbacks import CallbackPermanentError, CallbackRetryableError, URLPolicy
from compliance_service.contracts import canonical
from compliance_service.ledger import Ledger
from compliance_service.oceanbase import OceanBaseRepository
from compliance_service.service import InspectionService


class Cursor:
    description = []
    def __init__(self, row): self.row = row
    def execute(self, sql, params):
        self.sql, self.params = sql, params
    def fetchall(self): return [self.row]
    def close(self): pass


class Connection:
    def __init__(self, cursor): self.result = cursor
    def cursor(self): return self.result
    def close(self): pass


class Callback:
    timeout = 1
    policy = URLPolicy(allowed_hosts={'callbacks.example'})
    def __init__(self): self.sent = []; self.failure = None
    def send(self, url, payload, delivery_id):
        self.sent.append((url, canonical(payload), delivery_id))
        if self.failure: raise self.failure


class UpstreamIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.clock = [1000.0]
        self.ledger = Ledger(self.root/'service.sqlite3', clock=lambda: self.clock[0])
        document = Document()
        document.add_paragraph('2026-10-14 403105工作面未记录动力现象。未找到撤人记录。')
        buffer = io.BytesIO(); document.save(buffer); self.data = buffer.getvalue()
        self.cursor = Cursor({'mine_code':'001', 'file_name_id':'rock_ledger_1k8hqc4_rpt_wh5nwe',
                              'attachment_id':'file-a', 'attachment_name':'日报.docx',
                              'attachment_url':'https://files.example/daily.docx',
                              'report_date':'2026-10-14', 'report_time_year':2026})
        self.repository = OceanBaseRepository(lambda: Connection(self.cursor), environment='production')
        self.acquisition = AttachmentAcquisition(allowed_hosts={'files.example'}, fetcher=lambda _: io.BytesIO(self.data),
                                                environment='production')
        self.callback = Callback()
        self.service = InspectionService(self.ledger, repository=self.repository, acquisition=self.acquisition,
                                         callback_client=self.callback, environment='production', allowed_mines={'001'},
                                         status_map={'paused':'pause', 'resumed':'resume'})
        self.service.task_token = 'test-only-principal-token'
        self.client = TestClient(create_app(self.service))
        self.addCleanup(self.client.close)
        self.headers = {'X-Inspection-Token':self.service.task_token}
        self.body = {'mineId':'001','status':'running','checkCycle':'2026-10',
                     'items':[{'itemCode':'stop_work_unload'},{'itemCode':'relief_effects'}],
                     'resultCallbackUrl':'https://callbacks.example/result'}
        self.models = patch('compliance_service.config.load_model_settings', return_value={'ocr':{'enabled':False},'vision':{'enabled':False}})
        self.models.start(); self.addCleanup(self.models.stop)

    def submit(self, **changes):
        response = self.client.post('/inspection/tasks', json={**self.body, **changes}, headers=self.headers)
        self.assertEqual(200, response.status_code, response.text)
        self.assertEqual({'code':200,'success':True,'msg':'任务控制请求已生效'}, response.json())
        return response.headers['X-Inspection-Task']

    def task(self, key): return self.ledger.get_task(key, 'oceanbase', '001', 'production')

    def test_intake_two_items_real_parse_retry_same_payload_no_reexecution(self):
        key = self.submit()
        self.assertEqual(key, self.submit())
        self.assertEqual(2, len(self.task(key)['items']))
        self.assertFalse(any(r['result'] for r in self.task(key)['items']))
        for _ in range(2):
            result = self.service.run_once()
            self.assertNotEqual('failed', result['status'], result)
            self.assertNotEqual('no_pass', result['status'], result)
        for run in self.task(key)['items']:
            self.assertTrue(run['acquisition']['raw_rows'])
            self.assertTrue(run['snapshot']['files'][0]['content_hash'])
        self.callback.failure = CallbackRetryableError('lost acknowledgement')
        self.assertEqual('waiting', self.service.deliver_once()['state'])
        first = self.callback.sent[-1]
        self.clock[0] += 10
        self.callback.failure = None
        self.service.deliver_once(); self.service.deliver_once()
        retries = [p for p in self.callback.sent if p[2] == first[2]]
        self.assertEqual([first,first], retries)
        payload = json.loads(first[1])
        self.assertEqual('001', payload['mineId'])
        self.assertEqual('2026-10', payload['checkCycle'])
        self.assertIn('+00:00', payload['checkTime'])
        self.assertEqual('https://files.example/daily.docx',payload['attachments'][0]['attachmentUrl'])
        self.assertNotIn('attachmentContent',payload['attachments'][0])
        self.assertIsNone(self.service.run_once())
        self.assertEqual(key,self.submit())  # Terminal duplicate is not a new execution.
        self.assertIsNone(self.service.run_once())

    def test_pause_resume_same_runs_then_business_failure_is_visible(self):
        key = self.submit()
        before = {r['run_id'] for r in self.task(key)['items']}
        self.submit(status='paused')
        self.assertIsNone(self.service.run_once())
        self.submit(status='resumed')
        self.assertEqual(before,{r['run_id'] for r in self.task(key)['items']})
        self.service.run_once()
        self.callback.failure = CallbackPermanentError('business response unsuccessful')
        self.assertEqual('failed',self.service.deliver_once()['state'])
        self.assertIsNone(self.service.deliver_once())

    def test_auth_route_and_callback_snapshot_boundaries(self):
        self.assertEqual(404,self.client.post('/inspection/tasks',json=self.body).status_code)
        self.assertEqual(404,self.client.post('/inspection/tasks',json={**self.body,'mineId':'002'},headers=self.headers).status_code)
        key=self.submit()
        self.assertEqual(409,self.client.post('/inspection/tasks',json={**self.body,'resultCallbackUrl':'https://callbacks.example/new'},headers=self.headers).status_code)
        run=self.task(key)['items'][0]
        query={'source_system':'oceanbase','mine_code':'001','environment':'production'}
        url=f"/v1/runs/{run['run_id']}"
        self.assertEqual(404,self.client.get(url,params=query).status_code)
        self.assertEqual(200,self.client.get(url,params=query,headers=self.headers).status_code)
        for path in ('/debug','/v1/manual/runs','/v1/debug/bundles'):
            response=self.client.get(path) if path=='/debug' else self.client.post(path,json={})
            self.assertEqual(404,response.status_code)

    def test_confirmed_daily_issue_cannot_bypass_rule_preconditions(self):
        doc=Document()
        doc.add_paragraph('2026-10-14 403105工作面发生预警，明确未安排撤人。')
        buffer=io.BytesIO(); doc.save(buffer); self.data=buffer.getvalue()
        self.submit()
        result=self.service.run_once()
        self.assertTrue(result['daily_issues'])
        self.assertNotEqual('no_pass',result['status'],result)
        self.assertFalse(result['findings'])

    def test_running_pause_resumes_frozen_query_and_download(self):
        key=self.submit(items=[{'itemCode':'stop_work_unload'}])
        fetches=[]
        def fetch(url):
            fetches.append(url)
            self.ledger.control_task(key,'oceanbase','001','production','pause')
            return io.BytesIO(self.data)
        self.acquisition.fetcher=fetch
        with patch.object(self.repository,'query',wraps=self.repository.query) as query:
            paused=self.service.run_once()
            self.assertEqual('paused',paused['execution_state'])
            self.assertEqual('paused',self.task(key)['items'][0]['execution_state'])
            self.assertIsNone(self.service.run_once())
            self.submit(items=[{'itemCode':'stop_work_unload'}],status='resumed')
            self.assertNotEqual('failed',self.service.run_once()['status'])
            self.assertEqual(1,query.call_count)
        self.assertEqual(1,len(fetches))

    def test_failure_per_attachment_survives_callback_and_task_progress(self):
        self.acquisition.fetcher=lambda _: (_ for _ in ()).throw(TimeoutError('download timeout'))
        key=self.submit(items=[{'itemCode':'stop_work_unload'}])
        self.service.run_once()
        self.service.deliver_once()
        payload=json.loads(self.callback.sent[0][1])
        self.assertIn('file-a: attachment acquisition failed',payload['detail'])
        status=self.client.get('/inspection/tasks/'+key,params={'mine_id':'001'},headers=self.headers)
        self.assertEqual(200,status.status_code)
        self.assertEqual('completed',status.json()['state'])
        run=status.json()['items'][0]
        self.assertEqual('delivered',run['delivery_state'])
        events=self.ledger.get_progress(run['run_id'],'oceanbase','001','production')['events']
        self.assertTrue(any(e['stage']=='callback' and 'elapsed_ms' in e for e in events))


if __name__=='__main__': unittest.main()
