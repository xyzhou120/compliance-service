"""Binary document tests; no network, OCR or server is started."""

from hashlib import sha256
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject, NumberObject

from compliance_service.parsers import parse_document


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)

    def tearDown(self):
        self.directory.cleanup()

    def parse(self, name):
        return parse_document(self.root / name, file_id="file-1")

    def pdf(self, name="sample.pdf", *, text=True, blank=True, image=False, encrypted=False):
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        resources = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 30 250 Td (Face A1 value 0) Tj ET" if text else b"")
        if image:
            bitmap = DecodedStreamObject()
            bitmap.set_data(b"\xff\xff\xff")
            bitmap.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
                           NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1),
                           NameObject("/ColorSpace"): NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"): NumberObject(8)})
            resources[NameObject("/XObject")] = DictionaryObject({NameObject("/Im0"): writer._add_object(bitmap)})
            stream.set_data(stream.get_data() + b" q 100 0 0 100 0 0 cm /Im0 Do Q")
        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = writer._add_object(stream)
        if blank:
            writer.add_blank_page(width=300, height=300)
        if encrypted:
            writer.encrypt("secret")
        writer.write(self.root / name)

    def image(self):
        image = BytesIO()
        Image.new("RGB", (4, 4), "white").save(image, format="PNG")
        image.seek(0)
        return image

    def test_pdf_native_blank_page_and_stable_hash_locator(self):
        self.pdf()
        before = (self.root / "sample.pdf").read_bytes()
        first, second = self.parse("sample.pdf"), self.parse("sample.pdf")
        self.assertEqual(first, second)
        self.assertEqual(before, (self.root / "sample.pdf").read_bytes())
        self.assertEqual(first["content_hash"], sha256(before).hexdigest())
        self.assertIn("Face A1 value 0", first["blocks"][0]["text"])
        self.assertEqual(first["blocks"][0]["locator"], {"page": 1})
        self.assertEqual([c["state"] for c in first["coverage"]], ["readable", "empty"])
        self.pdf("changed.pdf", text=False)
        self.assertNotEqual(first["content_hash"], self.parse("changed.pdf")["content_hash"])

    def test_pdf_scan_and_blank_are_different(self):
        self.pdf(text=False, image=True)
        result = self.parse("sample.pdf")
        self.assertEqual(result["blocks"], [])
        self.assertEqual([c["state"] for c in result["coverage"]], ["unreadable", "empty"])
        self.assertIn("requires_ocr", result["issues"][0])

    def test_pdf_mixed_text_image_is_not_complete(self):
        self.pdf(image=True, blank=False)
        result = self.parse("sample.pdf")
        self.assertTrue(result["blocks"])
        self.assertEqual(result["coverage"][0]["state"], "partial")

    def test_pdf_unused_shared_image_resource_does_not_require_ocr(self):
        self.pdf(image=True, blank=False)
        from pypdf import PdfReader
        reader = PdfReader(self.root / "sample.pdf")
        writer = PdfWriter()
        page = writer.add_page(reader.pages[0])
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 30 250 Td (Original text) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(self.root / "shared.pdf")
        events = []
        result = parse_document(self.root / "shared.pdf", file_id="shared", progress_callback=events.append)
        self.assertEqual(result["coverage"][0]["state"], "readable")
        self.assertFalse(result["blocks"][0]["metadata"]["has_images"])
        self.assertEqual([e["event"] for e in events], ["page_started", "page_finished"])
        self.assertIn("elapsed_ms", events[-1])

    def test_pdf_progress_propagates_stale_lease(self):
        from compliance_service.contracts import StaleLease
        self.pdf(blank=False)
        def stale(event):
            raise StaleLease("expired")
        with self.assertRaises(StaleLease):
            parse_document(self.root / "sample.pdf", file_id="stale", progress_callback=stale)

    def test_pdf_inline_image_is_not_silently_covered_by_text(self):
        self.pdf(blank=False)
        from pypdf import PdfReader
        reader = PdfReader(self.root / "sample.pdf")
        writer = PdfWriter()
        page = writer.add_page(reader.pages[0])
        stream = DecodedStreamObject()
        stream.set_data(page.get_contents().get_data() + b" BI /W 1 /H 1 /CS /RGB /BPC 8 ID \xff\xff\xff EI\n")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(self.root / "inline.pdf")
        result = self.parse("inline.pdf")
        self.assertIn("Face A1", result["blocks"][0]["text"])
        self.assertEqual(result["coverage"][0]["state"], "partial")

    def test_pdf_vector_content_without_text_is_not_blank(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        stream = DecodedStreamObject()
        stream.set_data(b"10 10 200 200 re S\n")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(self.root / "vector.pdf")
        result = self.parse("vector.pdf")
        self.assertEqual(result["coverage"][0]["state"], "unreadable")
        self.assertEqual(result["coverage"][0]["reason"], "nonempty_page_without_extractable_text")

    def test_pdf_encrypted_and_corrupted_fail_explicitly(self):
        self.pdf(encrypted=True)
        self.assertEqual(self.parse("sample.pdf")["coverage"][0]["reason"], "encrypted_pdf")
        (self.root / "bad.pdf").write_bytes(b"%PDF-1.7\nbroken")
        self.assertEqual(self.parse("bad.pdf")["coverage"][-1]["state"], "failed")

    def test_docx_original_body_order_tables_and_stable_locators(self):
        doc = Document()
        doc.add_paragraph("  工作面 A1\t原文  ")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "工作面", "测值"
        table.cell(1, 0).text, table.cell(1, 1).text = "A1", "0"
        doc.add_paragraph("末段")
        doc.save(self.root / "sample.docx")
        result = self.parse("sample.docx")
        self.assertEqual([b["kind"] for b in result["blocks"]], ["paragraph", "table_row", "table_row", "paragraph"])
        self.assertEqual(result["blocks"][0]["text"], "  工作面 A1\t原文  ")
        self.assertEqual(result["blocks"][2]["metadata"]["cells"], ["A1", "0"])
        self.assertEqual(result["blocks"][2]["metadata"]["headers"], ["工作面", "测值"])
        self.assertEqual(result["blocks"][2]["locator"], {"table": 1, "row": 2})
        self.assertEqual(result, self.parse("sample.docx"))

    def test_docx_revision_image_and_header_are_explicit_gaps(self):
        doc = Document()
        paragraph = doc.add_paragraph("正文")
        revision = OxmlElement("w:del")
        run = OxmlElement("w:r")
        deleted = OxmlElement("w:delText")
        deleted.text = "删除的文字"
        run.append(deleted)
        revision.append(run)
        paragraph._p.append(revision)
        doc.add_picture(self.image())
        doc.sections[0].header.paragraphs[0].text = "页眉证据"
        doc.save(self.root / "sample.docx")
        result = self.parse("sample.docx")
        self.assertIn("删除的文字", result["blocks"][0]["text"])
        self.assertEqual(result["coverage"][0]["state"], "unreadable")
        self.assertTrue(any("embedded_object" in issue for issue in result["issues"]))
        self.assertTrue(any("non_body_content" in issue for issue in result["issues"]))

    def test_docx_merged_cell_does_not_invent_repeated_values(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "标题"
        table.cell(1, 0).text = "A1"
        doc.save(self.root / "sample.docx")
        result = self.parse("sample.docx")
        self.assertEqual(result["blocks"][0]["metadata"]["cells"], ["标题", ""])
        self.assertEqual(result["blocks"][1]["metadata"]["cells"], ["A1", ""])

    def workbook(self, name="sample.xlsx"):
        book = Workbook()
        sheet = book.active
        sheet.title = "测量"
        sheet.append(["工作面", "测值", "备注"])
        sheet.append(["A1", 0, " 原文 "])
        sheet.append([None, 2, "不应继承工作面"])
        sheet.append(["A2", "=1-1", None])
        sheet.merge_cells("A5:A6")
        sheet["A5"] = "合并工作面"
        sheet["B5"], sheet["B6"] = 3, 4
        sheet.row_dimensions[3].hidden = True
        sheet.column_dimensions.group("B", "C", hidden=True)
        hidden = book.create_sheet("隐藏表")
        hidden.sheet_state = "hidden"
        hidden.append(["原文", 0])
        book.create_sheet("空表")
        book.save(self.root / name)

    def test_xlsx_zero_formula_hidden_merge_and_no_forward_fill(self):
        self.workbook()
        result = self.parse("sample.xlsx")
        rows = [b for b in result["blocks"] if b["locator"]["sheet"] == "测量"]
        self.assertEqual(rows[1]["metadata"]["cells"], ["A1", "0", " 原文 "])
        self.assertEqual(rows[2]["metadata"]["cells"][0], "")
        self.assertTrue(rows[2]["metadata"]["row_hidden"])
        self.assertTrue(rows[1]["metadata"]["cell_details"][1]["column_hidden"])
        self.assertEqual(rows[3]["metadata"]["cell_details"][1]["formula"], "=1-1")
        self.assertIsNone(rows[3]["metadata"]["cell_details"][1]["value"])
        self.assertEqual(rows[5]["metadata"]["cells"][0], "")
        self.assertEqual(rows[5]["metadata"]["cell_details"][0]["merged_anchor"], "A5")
        self.assertTrue(any("formula_without_cached_value" in issue for issue in result["issues"]))
        self.assertTrue(any(b["metadata"]["sheet_state"] == "hidden" for b in result["blocks"]))
        self.assertEqual(next(c["state"] for c in result["coverage"] if c["unit_id"] == "sheet:空表"), "empty")
        self.assertEqual(result, self.parse("sample.xlsx"))

    def test_xlsx_cached_formula_zero_is_preserved(self):
        self.workbook()
        original = (self.root / "sample.xlsx").read_bytes()
        with ZipFile(BytesIO(original)) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        xml = entries["xl/worksheets/sheet1.xml"]
        self.assertIn(b"<f>1-1</f><v></v>", xml)
        entries["xl/worksheets/sheet1.xml"] = xml.replace(b"<f>1-1</f><v></v>", b"<f>1-1</f><v>0</v>")
        with ZipFile(self.root / "cached.xlsx", "w", ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        result = self.parse("cached.xlsx")
        row = next(b for b in result["blocks"] if b["locator"] == {"sheet": "测量", "row": 4})
        self.assertEqual(row["metadata"]["cells"][1], "0")
        self.assertEqual(row["metadata"]["cell_details"][1]["cached_value"], 0)
        self.assertFalse(result["issues"])

    def test_xlsx_error_and_drawing_are_explicit_gaps(self):
        from openpyxl.drawing.image import Image as SpreadsheetImage
        book = Workbook()
        sheet = book.active
        sheet.append(["工作面", "测值"])
        sheet.append(["A1", "#DIV/0!"])
        sheet.add_image(SpreadsheetImage(self.image()), "D1")
        book.save(self.root / "errors.xlsx")
        result = self.parse("errors.xlsx")
        self.assertTrue(any("cell_error:B2" in issue for issue in result["issues"]))
        self.assertTrue(any("drawing_requires_extraction" in issue for issue in result["issues"]))
        self.assertEqual(result["coverage"][-1]["state"], "unreadable")

    def test_unsupported_images_and_failure_do_not_raise(self):
        for name in ("old.xls", "old.doc", "other.txt"):
            (self.root / name).write_bytes(b"content")
            self.assertEqual(self.parse(name)["coverage"][0]["state"], "unsupported")
        (self.root / "scan.png").write_bytes(self.image().getvalue())
        self.assertEqual(self.parse("scan.png")["coverage"][0]["state"], "unreadable")
        for name in ("missing.docx", "bad.docx", "bad.xlsx"):
            if not name.startswith("missing"):
                (self.root / name).write_bytes(b"broken")
            self.assertEqual(self.parse(name)["coverage"][-1]["state"], "failed")


if __name__ == "__main__":
    unittest.main()
